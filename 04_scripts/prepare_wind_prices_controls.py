#!/usr/bin/env python3
"""
Clean the 2026-07-01 Wind terminal exports into thesis-ready price/control data.

The raw Wind CSVs have nine metadata rows followed by a shared date column and
many mixed-frequency series. This script preserves a full long-form extract,
deduplicates repeated series, and writes a monthly wide panel for 2019-01
through 2026-03.
"""

from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[1]
RAW_DIR = REPO / "03_data" / "raw" / "wind_exports_2026-07-01"
INTERIM_DIR = REPO / "03_data" / "interim"
PROCESSED_DIR = REPO / "03_data" / "processed"
TABLE_DIR = REPO / "03_data" / "output" / "tables"

SOURCE_FILES = [
    RAW_DIR / "wind_export_gallium_rotterdam_multi_series.csv",
    RAW_DIR / "wind_export_indium_rotterdam_controls_multi_series.csv",
    RAW_DIR / "wind_export_ev_steel_lead_semiconductors_controls.csv",
    RAW_DIR / "wind_export_minerals_data_2026-06-09.xlsx",
]

BLOOMBERG_FILES = [
    RAW_DIR / "bloomberg_export_mineral_prices_2026-07-01.xlsx",
]

THESIS_START = pd.Timestamp("2019-01-01")
THESIS_END = pd.Timestamp("2026-03-31")


SERIES_MAP: dict[str, dict[str, str]] = {
    "S5810680": dict(var="ga_int_rotterdam_9999_usd_kg", group="gallium", role="international_price", hs="811292"),
    "S5812124": dict(var="ga_china_avg_99_cny_kg", group="gallium", role="china_price", hs="811292"),
    "S6949818": dict(var="ga_china_industrial_9999_cny_kg", group="gallium", role="china_price", hs="811292"),
    "S6949819": dict(var="ga_china_high_purity_6n_cny_kg", group="gallium", role="china_price", hs="811292"),
    "S5811361": dict(var="ga_china_domestic_9999_cny_kg", group="gallium", role="china_price", hs="811292"),
    "S5810682": dict(var="geo2_int_rotterdam_9999_usd_kg", group="germanium", role="international_price", hs="282560"),
    "S5810681": dict(var="ge_metal_int_rotterdam_50ohm_usd_kg", group="germanium", role="international_price", hs="811292"),
    "S6949812": dict(var="ge_metal_china_avg_50ohm_cny_kg", group="germanium", role="china_price", hs="811292"),
    "S5801747": dict(var="ge_metal_china_shanghai_longtai_cny_kg", group="germanium", role="china_price", hs="811292"),
    "S5812121": dict(var="ge_metal_china_avg_50ohm_ebaiyin_cny_kg", group="germanium", role="china_price", hs="811292"),
    "S6949813": dict(var="geo2_china_99999_cny_kg", group="germanium", role="china_price", hs="282560"),
    "S6949753": dict(var="v_china_vn16_cny_10k_ton", group="vanadium", role="china_price", hs="811292;282530"),
    "S5705621": dict(var="v_china_ferrovanadium_50_cny_ton", group="vanadium", role="china_price", hs="811292;282530"),
    "S9976462": dict(var="v_china_v2o5_98_flake_cny_ton", group="vanadium", role="china_price", hs="282530"),
    "S5810701": dict(var="v_int_ferrovanadium_80_rotterdam_usd_kg", group="vanadium", role="international_price", hs="811292;282530"),
    "S5810702": dict(var="v_int_v2o5_98_rotterdam_usd_lb", group="vanadium", role="international_price", hs="282530"),
    "S5810683": dict(var="in_int_rotterdam_9999_usd_kg", group="indium", role="international_price", hs="811292"),
    "S5812125": dict(var="in_china_avg_9999_cny_kg", group="indium", role="china_price", hs="811292"),
    "S6949800": dict(var="in_china_crude_99_cny_kg", group="indium", role="china_price", hs="811292"),
    "S6949799": dict(var="in_china_refined_99995_cny_kg", group="indium", role="china_price", hs="811292"),
    "S6949801": dict(var="in_china_fine_99_usd_kg", group="indium", role="china_price", hs="811292"),
    "S5810671": dict(var="sb_int_rotterdam_9965_usd_ton", group="antimony", role="international_price", hs="8110"),
    "S0048100": dict(var="sb_china_antimony_1_nonferrous_cny_ton", group="antimony", role="china_price", hs="8110"),
    "S6949788": dict(var="sb_china_ingot_9965_cny_ton", group="antimony", role="china_price", hs="8110"),
    "S0048101": dict(var="sb_china_antimony_2_nonferrous_cny_ton", group="antimony", role="china_price", hs="8110"),
    "S6949789": dict(var="sb_china_concentrate_50_cny_metal_ton", group="antimony", role="china_price", hs="261710"),
    "S6949790": dict(var="sb_china_oxide_995_cny_ton", group="antimony", role="china_price", hs="282580"),
    "S6949791": dict(var="sb_china_ingot_9965_fob_usd_ton", group="antimony", role="export_price", hs="8110"),
    "S5811342": dict(var="sb_china_ingot_0_9990_cny_ton", group="antimony", role="china_price", hs="8110"),
    "S6949792": dict(var="sb_china_oxide_995_usd_ton", group="antimony", role="export_or_usd_price", hs="282580"),
    "S5810700": dict(var="w_int_ferrotungsten_80_85_rotterdam_usd_kg", group="tungsten", role="international_price", hs="8101;284180"),
    "S5705619": dict(var="w_china_concentrate_65_cny_ton", group="tungsten", role="china_price", hs="284180;8101"),
    "S5801732": dict(var="w_china_rod_1_9995_cny_kg", group="tungsten", role="china_price", hs="8101"),
    "S5811355": dict(var="w_china_powder_997_cny_kg", group="tungsten", role="china_price", hs="8101"),
    "S5811353": dict(var="w_china_carbide_powder_997_cny_kg", group="tungsten", role="china_price", hs="8101"),
    "S5812131": dict(var="w_china_avg_1_9995_cny_kg", group="tungsten", role="china_price", hs="8101"),
    "S5810673": dict(var="bi_int_rotterdam_9999_usd_lb", group="bismuth", role="international_price", hs="8106"),
    "S6949806": dict(var="bi_china_refined_9999_cny_ton", group="bismuth", role="china_price", hs="8106"),
    "S5812122": dict(var="bi_china_avg_9999_cny_ton", group="bismuth", role="china_price", hs="8106"),
    "S6949807": dict(var="bi_china_oxide_999_cny_ton", group="bismuth", role="china_price", hs="8106"),
    "S5812126": dict(var="te_china_avg_9999_cny_kg", group="tellurium", role="china_price", hs="280450"),
    "S6949834": dict(var="te_china_ingot_9999_cny_kg", group="tellurium", role="china_price", hs="280450"),
    "S5801746": dict(var="te_china_domestic_9999_cny_kg", group="tellurium", role="china_price", hs="280450"),
    "S5801744": dict(var="mo_china_metal_1_9995_cny_kg", group="molybdenum", role="china_price", hs="284170;282570"),
    "S6949736": dict(var="mo_china_concentrate_40_45_cny_ton_unit", group="molybdenum", role="china_price", hs="284170;282570"),
    "S5705615": dict(var="mo_china_oxide_gt51_cny_ton_unit", group="molybdenum", role="china_price", hs="282570"),
    "S6949738": dict(var="mo_china_industrial_oxide_51_cny_ton_unit", group="molybdenum", role="china_price", hs="282570"),
    "S0243695": dict(var="mo_us_fcx_avg_selling_price_usd_lb", group="molybdenum", role="international_proxy", hs="284170;282570"),
    "S6949744": dict(var="mo_china_powder_first_grade_cny_10k_ton", group="molybdenum", role="china_price", hs="284170;282570"),
    "S6949743": dict(var="mo_china_high_purity_moo3_cny_10k_ton", group="molybdenum", role="china_price", hs="282570"),
    "S6500784": dict(var="demand_global_semiconductor_sales_3mma_usd_bn", group="demand", role="demand_control", hs=""),
    "M0041963": dict(var="macro_china_total_industrial_added_value_cny_100mn", group="macro", role="macro_control", hs=""),
    "M0017126": dict(var="macro_china_manufacturing_pmi_pct", group="macro", role="macro_control", hs=""),
    "M0001318": dict(var="macro_china_ppi_nonmetal_ores_yoy_pct", group="macro", role="macro_control", hs=""),
    "M0096797": dict(var="macro_china_ppi_nonmetal_ores_mom_pct", group="macro", role="macro_control", hs=""),
    "M0000200": dict(var="fx_eur_usd_ice_spot", group="fx", role="fx_control", hs=""),
    "M0061688": dict(var="fx_usd_cny_wind_consensus", group="fx", role="fx_control", hs=""),
    "M0067855": dict(var="fx_usd_cny_cfets_spot", group="fx", role="fx_control", hs=""),
    "M0096812": dict(var="macro_china_ppi_chemicals_mom_pct", group="macro", role="macro_control", hs=""),
    "M0001333": dict(var="macro_china_ppi_chemicals_yoy_pct", group="macro", role="macro_control", hs=""),
    "M0096796": dict(var="macro_china_ppi_nonferrous_ores_mom_pct", group="macro", role="macro_control", hs=""),
    "M0001317": dict(var="macro_china_ppi_nonferrous_ores_yoy_pct", group="macro", role="macro_control", hs=""),
    "S6137468": dict(var="demand_china_power_battery_loading_mwh", group="demand", role="demand_control", hs="2504;282520;8105"),
    "S6139213": dict(var="demand_china_nev_bev_sales_units", group="demand", role="demand_control", hs="2504;282520;8105"),
    "S6131303": dict(var="demand_china_nev_passenger_bev_output_units", group="demand", role="demand_control", hs="2504;282520;8105"),
    "S6131306": dict(var="demand_china_nev_commercial_bev_output_units", group="demand", role="demand_control", hs="2504;282520;8105"),
    "S0027374": dict(var="demand_china_crude_steel_output_10k_tons", group="demand", role="demand_control", hs="811292;282530;284180;8101;284170;282570"),
    "S0070265": dict(var="demand_china_lead_acid_battery_output_kvah", group="demand", role="demand_control", hs="8110;261710;282580"),
    "S5603672": dict(var="demand_china_semiconductor_chip_output_10k_units", group="demand", role="demand_control", hs="811292"),
    "S5621787": dict(var="demand_china_semiconductor_sales_value_usd_100mn", group="demand", role="demand_control", hs="811292"),
    "K1272331": dict(var="graphite_china_natural_flake_cny_ton", group="graphite", role="china_price", hs="2504"),
    "S9980000": dict(var="graphite_china_shandong_natural_spherical_cny_ton", group="graphite", role="china_price", hs="3801"),
    "TEEUUQPU Index": dict(var="te_int_bloomberg_teeuuqpu_px_last", group="tellurium", role="international_price_assumed", hs="280450"),
    "GPCNLCLW Index": dict(var="graphite_int_bloomberg_gpcnlclw_px_last", group="graphite", role="international_price_unverified", hs="2504;3801"),
    "GPCNASPZ Index": dict(var="graphite_int_bloomberg_gpcnaspz_px_last", group="graphite", role="international_price_unverified", hs="2504;3801"),
    "GPCNASKZ Index": dict(var="graphite_int_bloomberg_gpcnaskz_px_last", group="graphite", role="international_price_unverified", hs="2504;3801"),
    "GPCNN190 Index": dict(var="graphite_int_bloomberg_gpcnn190_px_last", group="graphite", role="international_price_unverified", hs="2504;3801"),
    "GPCNP195 Index": dict(var="graphite_int_bloomberg_gpcnp195_px_last", group="graphite", role="international_price_unverified", hs="2504;3801"),
    "VNEUJOQT Index": dict(var="v_int_bloomberg_vneujoqt_px_last", group="vanadium", role="international_price_unverified", hs="282530;811292"),
    "TGEUWOQR Index": dict(var="w_int_bloomberg_tgeuwoqr_px_last", group="tungsten", role="international_price_assumed", hs="284180;8101"),
}

VARIABLE_ORDER = [meta["var"] for meta in SERIES_MAP.values()]


PRODUCT_AUDIT = [
    ("2504", "Graphite", "China natural flake graphite is present through 2026-03. Bloomberg graphite tickers are present, but their descriptions/grades were not preserved in the workbook.", "Verify Bloomberg graphite ticker descriptions; then choose the natural flake 94-95% / -100 mesh series as primary."),
    ("811292", "Ga/Ge/In/V bundle", "Gallium, germanium metal, indium, and vanadium prices are present with China and international series. Global and China semiconductor controls are present.", "No niobium price."),
    ("282560", "GeO2/ZrO2", "GeO2 has Rotterdam and China series.", "No zirconium dioxide/sponge price."),
    ("8110", "Antimony metal", "Rotterdam antimony, several China antimony metal series, and lead-acid battery demand are present.", "Good enough for Sb metal wedge."),
    ("261710", "Antimony ores/concentrates", "China antimony concentrate 50% is present.", "No international concentrate benchmark."),
    ("282580", "Antimony oxides", "China antimony oxide in CNY and USD is present.", "No Rotterdam/Europe antimony trioxide benchmark."),
    ("8106", "Bismuth", "Rotterdam bismuth, China refined bismuth, average bismuth, and bismuth oxide are present.", "Bismuth oxide international benchmark still missing."),
    ("284180", "Tungstates/APT", "Tungsten China series and international `TGEUWOQR Index` are present; treat tungsten as usable under the current assumption.", "Archive `TGEUWOQR Index` description tomorrow, but no new tungsten download is required unless the description contradicts the assumption."),
    ("8101", "Tungsten articles", "Tungsten rod, powder, carbide powder, average tungsten, Rotterdam ferro-tungsten, and `TGEUWOQR Index` are present.", "No new tungsten download required under the current assumption."),
    ("280450", "Tellurium/boron", "China tellurium series and international `TEEUUQPU Index` are present; treat tellurium as usable under the current assumption.", "Archive `TEEUUQPU Index` description tomorrow; boron contamination proxy absent."),
    ("284170", "Molybdates", "China molybdenum concentrate/oxide/metal/powder series and annual FCX selling price are present.", "Candidate-control/proxy family; no monthly Rotterdam/Asia molybdenum oxide benchmark."),
    ("282570", "Molybdenum oxides", "China molybdenum oxide and high-purity MoO3 series are present.", "Candidate-control/proxy family; no monthly international molybdenum oxide benchmark."),
    ("284690", "Non-cerium rare-earth compounds", "No rare-earth series in these files.", "Drop rare-earth price-wedge analysis unless new data are found."),
    ("284610", "Cerium compounds", "No cerium series in these files.", "No price wedge."),
    ("2818", "Alumina", "No alumina series in these files.", "Missing."),
    ("8111", "Manganese articles", "No manganese series in these files.", "Missing."),
    ("2811", "Inorganic acids", "Chemical PPI controls are present, but no acid price.", "Acid-specific price missing."),
    ("2615", "Nb/Ta/V/Zr ores", "Vanadium and molybdenum proxies are present.", "Niobium, tantalum, and zirconium ore proxies missing."),
    ("811221", "Chromium", "No chromium/ferrochrome series in these files.", "Missing."),
    ("380130", "Carbon electrode pastes", "No electrode paste, needle coke, or synthetic graphite series in these files.", "Missing."),
    ("282530", "Vanadium oxides", "Rotterdam V2O5, China V2O5 flake, China V nitrogen/ferrovanadium, and China crude steel demand are present.", "China V2O5 starts in 2020-10, so use ferrovanadium/V nitrogen for a full-window robustness proxy."),
    ("282520", "Lithium oxide/hydroxide", "Power-battery and BEV demand controls are present.", "Lithium hydroxide/carbonate prices missing."),
    ("8103", "Tantalum", "No tantalum series in these files.", "Missing."),
    ("8108", "Titanium", "No titanium series in these files.", "Missing."),
    ("8109", "Zirconium", "No zirconium series in these files.", "Missing."),
    ("8104", "Magnesium", "No magnesium series in these files.", "Missing."),
    ("8105", "Cobalt", "Power-battery and BEV demand controls are present.", "Cobalt price series missing in these files."),
    ("282110", "Iron oxide pigments", "No iron oxide pigment series in these files.", "Missing."),
    ("2820", "Manganese oxides", "No manganese oxide/EMD series in these files.", "Missing."),
    ("2833", "Sulfates", "Broad chemical PPI controls are present.", "Sulfate-specific prices missing."),
]

PRICE_WEDGE_AUDIT = [
    ("Gallium", "811292", "ga_int_rotterdam_9999_usd_kg", "ga_china_domestic_9999_cny_kg; ga_china_industrial_9999_cny_kg; ga_china_avg_99_cny_kg", "complete", "Use Rotterdam gallium against China domestic/industrial gallium."),
    ("Germanium metal", "811292", "ge_metal_int_rotterdam_50ohm_usd_kg", "ge_metal_china_avg_50ohm_cny_kg; ge_metal_china_shanghai_longtai_cny_kg; ge_metal_china_avg_50ohm_ebaiyin_cny_kg", "complete", "Use Rotterdam metal germanium against China 50 ohm/cm series."),
    ("Indium", "811292", "in_int_rotterdam_9999_usd_kg", "in_china_refined_99995_cny_kg; in_china_avg_9999_cny_kg; in_china_crude_99_cny_kg", "within_bundle_control_complete", "Indium metal is complete as a within-bundle comparison/proxy; the 2025 indium controls cover indium phosphide and organo-indium compounds, not ordinary indium metal."),
    ("Antimony metal", "8110", "sb_int_rotterdam_9965_usd_ton", "sb_china_antimony_1_nonferrous_cny_ton; sb_china_ingot_9965_cny_ton; sb_china_ingot_0_9990_cny_ton", "complete", "Use Rotterdam 99.65% antimony against China ingot/non-ferrous benchmarks."),
    ("Bismuth", "8106", "bi_int_rotterdam_9999_usd_lb", "bi_china_refined_9999_cny_ton; bi_china_avg_9999_cny_ton", "complete", "Use Rotterdam bismuth against China refined/average bismuth."),
    ("GeO2", "282560", "geo2_int_rotterdam_9999_usd_kg", "geo2_china_99999_cny_kg", "usable_partial", "International GeO2 is present but ends before 2026-03 in this export; domestic is complete."),
    ("Vanadium", "282530;811292", "v_int_v2o5_98_rotterdam_usd_lb; v_int_ferrovanadium_80_rotterdam_usd_kg; v_int_bloomberg_vneujoqt_px_last", "v_china_v2o5_98_flake_cny_ton; v_china_vn16_cny_10k_ton; v_china_ferrovanadium_50_cny_ton", "complete_from_2020m10_with_full_proxy", "Exact China V2O5 is now present from 2020-10; V nitrogen/ferrovanadium give full-window proxies."),
    ("Tungsten", "284180;8101", "w_int_ferrotungsten_80_85_rotterdam_usd_kg; w_int_bloomberg_tgeuwoqr_px_last", "w_china_concentrate_65_cny_ton; w_china_rod_1_9995_cny_kg; w_china_powder_997_cny_kg; w_china_carbide_powder_997_cny_kg; w_china_avg_1_9995_cny_kg", "complete_assumed", "Treat `TGEUWOQR Index` as the intended international tungsten benchmark under the current assumption; archive DES tomorrow."),
    ("Molybdenum", "284170;282570", "mo_us_fcx_avg_selling_price_usd_lb", "mo_china_oxide_gt51_cny_ton_unit; mo_china_industrial_oxide_51_cny_ton_unit; mo_china_metal_1_9995_cny_kg", "candidate_control_proxy_optional", "China molybdenum coverage is strong; only annual FCX international proxy is present. A monthly molybdenum oxide or ferro-moly benchmark is useful only as a candidate-control/proxy series, not as a treated-product wedge."),
    ("Antimony oxides", "282580", "none", "sb_china_oxide_995_cny_ton; sb_china_oxide_995_usd_ton", "download_required", "China antimony oxide is present; need Rotterdam/Europe/CIF antimony trioxide."),
    ("Tellurium", "280450", "te_int_bloomberg_teeuuqpu_px_last", "te_china_avg_9999_cny_kg; te_china_ingot_9999_cny_kg; te_china_domestic_9999_cny_kg", "complete_assumed", "Treat `TEEUUQPU Index` as the intended international tellurium benchmark under the current assumption; archive DES tomorrow."),
    ("Graphite", "2504;3801", "graphite_int_bloomberg_gpcnlclw_px_last; graphite_int_bloomberg_gpcnaspz_px_last; graphite_int_bloomberg_gpcnaskz_px_last; graphite_int_bloomberg_gpcnn190_px_last; graphite_int_bloomberg_gpcnp195_px_last", "graphite_china_natural_flake_cny_ton; graphite_china_shandong_natural_spherical_cny_ton; legacy graphite_flake_avg in 03_data/raw/bloomberg_prices_monthly.csv through 2024-12", "present_needs_bloomberg_description", "Bloomberg graphite tickers are present but descriptions were not exported; verify grade/location before using."),
]


def read_export_rows(path: Path) -> list[list[object]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="gb18030", newline="") as f:
            return list(csv.reader(f))
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=False, data_only=True)
        worksheet = workbook.worksheets[0]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    raise ValueError(f"Unsupported export type: {path}")


def parse_wind_export(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = read_export_rows(path)

    first_data_idx = None
    for idx, row in enumerate(rows):
        if row and pd.notna(pd.to_datetime(row[0], errors="coerce")):
            first_data_idx = idx
            break
    if first_data_idx is None:
        raise ValueError(f"Could not find first date row in {path}")

    meta_keys = [str(row[0]) for row in rows[:first_data_idx]]
    n_cols = max(len(row) for row in rows[:first_data_idx])
    inventory_rows = []
    value_frames = []

    parsed_dates = []
    data_rows = []
    for row in rows[first_data_idx:]:
        if not row:
            continue
        date = pd.to_datetime(row[0], errors="coerce")
        if pd.isna(date):
            continue
        parsed_dates.append(date)
        if len(row) < n_cols:
            row = row + [""] * (n_cols - len(row))
        data_rows.append(row)

    dates = pd.Series(parsed_dates, name="date")

    for pos in range(1, n_cols):
        md = {
            meta_keys[i]: rows[i][pos] if pos < len(rows[i]) and rows[i][pos] is not None else ""
            for i in range(len(meta_keys))
        }
        series_id = str(md.get("ID", ""))
        mapped = SERIES_MAP.get(
            series_id,
            dict(var=f"unmapped_{series_id or pos}", group="unmapped", role="unmapped", hs=""),
        )
        raw_values = [row[pos] if pos < len(row) else "" for row in data_rows]
        values = pd.to_numeric(
            pd.Series(raw_values).map(
                lambda val: str(val).replace(",", "") if val not in (None, "") else None
            ),
            errors="coerce",
        )
        has_value = values.notna()
        in_window = (dates >= THESIS_START) & (dates <= THESIS_END) & has_value
        inventory_rows.append(
            {
                "raw_file": path.name,
                "raw_position": pos,
                "series_id": series_id,
                "canonical_variable": mapped["var"],
                "group": mapped["group"],
                "role": mapped["role"],
                "hs_mapping": mapped["hs"],
                "country": md.get("Country", ""),
                "raw_name": md.get("Name", ""),
                "frequency": md.get("Frequency", ""),
                "unit": md.get("Unit", ""),
                "time_period_metadata": md.get("Time Period", ""),
                "source": md.get("Source", ""),
                "update_date": md.get("Update", ""),
                "observations_full": int(has_value.sum()),
                "first_observation": dates[has_value].min().date().isoformat() if has_value.any() else "",
                "last_observation": dates[has_value].max().date().isoformat() if has_value.any() else "",
                "observations_thesis_window": int(in_window.sum()),
                "first_thesis_observation": dates[in_window].min().date().isoformat() if in_window.any() else "",
                "last_thesis_observation": dates[in_window].max().date().isoformat() if in_window.any() else "",
            }
        )
        frame = pd.DataFrame(
            {
                "date": dates,
                "value": values,
                "raw_file": path.name,
                "raw_position": pos,
                "series_id": series_id,
                "canonical_variable": mapped["var"],
                "group": mapped["group"],
                "role": mapped["role"],
                "hs_mapping": mapped["hs"],
                "country": md.get("Country", ""),
                "raw_name": md.get("Name", ""),
                "frequency": md.get("Frequency", ""),
                "unit": md.get("Unit", ""),
                "source": md.get("Source", ""),
                "update_date": md.get("Update", ""),
            }
        )
        value_frames.append(frame[frame["value"].notna()])

    return pd.DataFrame(inventory_rows), pd.concat(value_frames, ignore_index=True)


def parse_bloomberg_bdh_export(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = read_export_rows(path)
    tickers = [str(value) if value is not None else "" for value in rows[3][1:]]
    fields = [str(value) if value is not None else "" for value in rows[5][1:]]
    inventory_rows = []
    value_frames = []

    parsed_dates = []
    data_rows = []
    for row in rows[6:]:
        if not row:
            continue
        date = pd.to_datetime(row[0], errors="coerce")
        if pd.isna(date):
            continue
        parsed_dates.append(date)
        data_rows.append(row)
    dates = pd.Series(parsed_dates, name="date")

    for offset, ticker in enumerate(tickers, start=1):
        if not ticker:
            continue
        field = fields[offset - 1] if offset - 1 < len(fields) else ""
        mapped = SERIES_MAP.get(
            ticker,
            dict(var=f"unmapped_{ticker.replace(' ', '_').lower()}", group="unmapped", role="unmapped", hs=""),
        )
        raw_values = [row[offset] if offset < len(row) else "" for row in data_rows]
        values = pd.to_numeric(
            pd.Series(raw_values).map(
                lambda val: str(val).replace(",", "") if val not in (None, "", "#N/A N/A") else None
            ),
            errors="coerce",
        )
        has_value = values.notna()
        in_window = (dates >= THESIS_START) & (dates <= THESIS_END) & has_value
        inventory_rows.append(
            {
                "raw_file": path.name,
                "raw_position": offset,
                "series_id": ticker,
                "canonical_variable": mapped["var"],
                "group": mapped["group"],
                "role": mapped["role"],
                "hs_mapping": mapped["hs"],
                "country": "",
                "raw_name": f"Bloomberg BDH {ticker} {field}; description not exported",
                "frequency": "Every Day",
                "unit": "",
                "time_period_metadata": f"{rows[0][1]}:{rows[1][1]}",
                "source": "Bloomberg",
                "update_date": "",
                "observations_full": int(has_value.sum()),
                "first_observation": dates[has_value].min().date().isoformat() if has_value.any() else "",
                "last_observation": dates[has_value].max().date().isoformat() if has_value.any() else "",
                "observations_thesis_window": int(in_window.sum()),
                "first_thesis_observation": dates[in_window].min().date().isoformat() if in_window.any() else "",
                "last_thesis_observation": dates[in_window].max().date().isoformat() if in_window.any() else "",
            }
        )
        frame = pd.DataFrame(
            {
                "date": dates,
                "value": values,
                "raw_file": path.name,
                "raw_position": offset,
                "series_id": ticker,
                "canonical_variable": mapped["var"],
                "group": mapped["group"],
                "role": mapped["role"],
                "hs_mapping": mapped["hs"],
                "country": "",
                "raw_name": f"Bloomberg BDH {ticker} {field}; description not exported",
                "frequency": "Every Day",
                "unit": "",
                "source": "Bloomberg",
                "update_date": "",
            }
        )
        value_frames.append(frame[frame["value"].notna()])

    return pd.DataFrame(inventory_rows), pd.concat(value_frames, ignore_index=True)


def choose_preferred_duplicate_rows(long: pd.DataFrame) -> pd.DataFrame:
    long = long.copy()
    long["preferred_file"] = False
    is_indium_or_control = long["group"].isin(["indium", "demand", "macro", "fx"])
    long.loc[
        is_indium_or_control
        & long["raw_file"].eq("wind_export_indium_rotterdam_controls_multi_series.csv"),
        "preferred_file",
    ] = True
    long.loc[
        ~is_indium_or_control
        & long["raw_file"].eq("wind_export_gallium_rotterdam_multi_series.csv"),
        "preferred_file",
    ] = True
    long = long.sort_values(
        ["canonical_variable", "date", "preferred_file", "raw_position"],
        ascending=[True, True, False, True],
    )
    return long.drop_duplicates(["canonical_variable", "date"], keep="first")


def write_product_audit(path: Path) -> pd.DataFrame:
    df = pd.DataFrame(PRODUCT_AUDIT, columns=["hs", "product", "coverage_now", "remaining_gap"])
    df.to_csv(path, index=False)
    return df


def _series_group_summary(wide: pd.DataFrame, variables: str) -> tuple[int, str, str]:
    if variables == "none":
        return 0, "", ""
    if variables.startswith("legacy graphite_flake_avg"):
        legacy_path = REPO / "03_data" / "raw" / "bloomberg_prices_monthly.csv"
        if not legacy_path.exists():
            return 0, "", ""
        legacy = pd.read_csv(legacy_path, parse_dates=["ym"])
        sub = legacy[legacy["ym"].between(THESIS_START, THESIS_END, inclusive="both")]
        usable = sub["graphite_flake_avg"].notna()
        if not usable.any():
            return 0, "", ""
        return int(usable.sum()), sub.loc[usable, "ym"].min().date().isoformat(), sub.loc[usable, "ym"].max().date().isoformat()

    vars_list = [var.strip() for var in variables.split(";") if var.strip()]
    best_count = 0
    first = ""
    last = ""
    for var in vars_list:
        if var not in wide.columns:
            continue
        usable = wide[var].notna()
        if int(usable.sum()) > best_count:
            best_count = int(usable.sum())
            first = wide.loc[usable, "month_date"].min()
            last = wide.loc[usable, "month_date"].max()
            first = pd.to_datetime(first).date().isoformat() if pd.notna(first) else ""
            last = pd.to_datetime(last).date().isoformat() if pd.notna(last) else ""
    return best_count, first, last


def write_price_wedge_audit(path: Path, wide: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for candidate, hs, international, domestic, status, notes in PRICE_WEDGE_AUDIT:
        int_months, int_first, int_last = _series_group_summary(wide, international)
        dom_months, dom_first, dom_last = _series_group_summary(wide, domestic)
        rows.append(
            {
                "candidate": candidate,
                "hs_mapping": hs,
                "status": status,
                "international_series": international,
                "international_months": int_months,
                "international_first_month": int_first,
                "international_last_month": int_last,
                "domestic_series": domestic,
                "domestic_months": dom_months,
                "domestic_first_month": dom_first,
                "domestic_last_month": dom_last,
                "notes": notes,
            }
        )
    df = pd.DataFrame(rows)
    df.to_csv(path, index=False)
    return df


def write_markdown_audit(
    path: Path,
    inventory: pd.DataFrame,
    product_audit: pd.DataFrame,
    price_wedge_audit: pd.DataFrame,
    wide: pd.DataFrame,
) -> None:
    blank = inventory[inventory["observations_thesis_window"].eq(0)]
    duplicate_count = int(inventory.duplicated("series_id", keep=False).sum())
    lines = [
        "# Wind Price and Control Export Audit, 2026-07-01",
        "",
        "## Files Parsed",
        "",
    ]
    for source in SOURCE_FILES:
        lines.append(f"- `{source.relative_to(REPO)}`")
    for source in BLOOMBERG_FILES:
        lines.append(f"- `{source.relative_to(REPO)}`")
    lines += [
        "",
        "## Clean Outputs",
        "",
        "- `03_data/processed/wind_prices_controls_monthly_2019_2026m03.csv`",
        "- `03_data/interim/wind_prices_controls_long_2026-07-01.csv`",
        "- `03_data/output/tables/wind_series_inventory_2026-07-01.csv`",
        "- `03_data/output/tables/wind_product_coverage_audit_2026-07-01.csv`",
        "- `03_data/output/tables/price_wedge_candidate_coverage_2026-07-01.csv`",
        "",
        "## Summary",
        "",
        f"- Raw series columns parsed: {len(inventory)}.",
        f"- Unique Wind series IDs: {inventory['series_id'].nunique()}.",
        f"- Duplicate raw series columns: {duplicate_count}. These are deduplicated in the clean panel.",
        f"- Monthly panel shape: {wide.shape[0]} months x {wide.shape[1] - 3} data columns.",
        "- Daily and weekly prices are averaged to monthly means; monthly series keep their monthly observations.",
        "- The panel is trimmed to January 2019 through March 2026.",
        "",
        "## Important Caveats",
        "",
        "- The later CSV-exported China V2O5 flake series is blank, but the older Wind workbook supplies usable China V2O5 from 2020-10 onward.",
        "- Bloomberg graphite tickers are present, but the workbook did not preserve descriptions, grades, locations, or units.",
        "- The exports contain no rare-earth price series; rare-earth price-wedge analysis should be dropped unless new data are found.",
        "- Tellurium is treated as complete under the current assumption that `TEEUUQPU Index` is the intended international tellurium benchmark; archive the DES tomorrow.",
        "- Tungsten is treated as complete/usable under the current assumption that `TGEUWOQR Index` is the intended international tungsten benchmark; archive the DES tomorrow.",
        "- Remaining treated/partial-treated price download: international antimony oxide.",
        "- International molybdenum is useful only as a candidate-control/proxy benchmark, not as a treated-product wedge.",
        "- The new controls file adds BEV sales/output, crude steel, lead-acid battery output, and China semiconductor sales value.",
        "",
        "## Blank Or Superseded Raw Export Columns",
        "",
    ]
    if blank.empty:
        lines.append("- None.")
    else:
        for _, row in blank.iterrows():
            lines.append(
                f"- `{row['raw_file']}`: `{row['series_id']}` / "
                f"`{row['canonical_variable']}`: {row['raw_name']}"
            )
    lines += [
        "",
        "## Price-Wedge Candidate Coverage",
        "",
        "| Candidate | Status | International months | Domestic months | Notes |",
        "|---|---|---:|---:|---|",
    ]
    for _, row in price_wedge_audit.iterrows():
        lines.append(
            f"| {row['candidate']} | {row['status']} | "
            f"{row['international_months']} | {row['domestic_months']} | {row['notes']} |"
        )
    lines += [
        "",
        "## Product Coverage",
        "",
        "| HS | Product | Coverage now | Remaining gap |",
        "|---|---|---|---|",
    ]
    for _, row in product_audit.iterrows():
        lines.append(f"| {row['hs']} | {row['product']} | {row['coverage_now']} | {row['remaining_gap']} |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    for directory in [INTERIM_DIR, PROCESSED_DIR, TABLE_DIR]:
        directory.mkdir(parents=True, exist_ok=True)

    inventories = []
    longs = []
    for source in SOURCE_FILES:
        inv, long = parse_wind_export(source)
        inventories.append(inv)
        longs.append(long)
    for source in BLOOMBERG_FILES:
        inv, long = parse_bloomberg_bdh_export(source)
        inventories.append(inv)
        longs.append(long)

    inventory = pd.concat(inventories, ignore_index=True)
    raw_long = pd.concat(longs, ignore_index=True)
    clean_long = choose_preferred_duplicate_rows(raw_long)

    monthly_source = clean_long[
        clean_long["date"].between(THESIS_START, THESIS_END, inclusive="both")
    ].copy()
    monthly_source["month_date"] = monthly_source["date"].dt.to_period("M").dt.to_timestamp()
    monthly = (
        monthly_source.groupby(["month_date", "canonical_variable"], as_index=False)["value"]
        .mean()
        .pivot(index="month_date", columns="canonical_variable", values="value")
    )
    all_months = pd.date_range("2019-01-01", "2026-03-01", freq="MS")
    monthly = monthly.reindex(all_months)
    ordered_cols = [col for col in VARIABLE_ORDER if col in monthly.columns and monthly[col].notna().any()]
    ordered_cols += [col for col in monthly.columns if col not in ordered_cols]
    monthly = monthly[ordered_cols].reset_index(names="month_date")
    monthly.insert(1, "year", monthly["month_date"].dt.year)
    monthly.insert(2, "month", monthly["month_date"].dt.month)

    inventory["duplicate_series_id"] = inventory.duplicated("series_id", keep=False)
    inventory["usable_in_thesis_window"] = inventory["observations_thesis_window"] > 0
    product_audit = write_product_audit(TABLE_DIR / "wind_product_coverage_audit_2026-07-01.csv")
    price_wedge_audit = write_price_wedge_audit(
        TABLE_DIR / "price_wedge_candidate_coverage_2026-07-01.csv",
        monthly,
    )

    inventory.to_csv(TABLE_DIR / "wind_series_inventory_2026-07-01.csv", index=False)
    clean_long.to_csv(INTERIM_DIR / "wind_prices_controls_long_2026-07-01.csv", index=False)
    monthly.to_csv(PROCESSED_DIR / "wind_prices_controls_monthly_2019_2026m03.csv", index=False)
    write_markdown_audit(
        REPO / "03_data" / "wind_price_control_export_audit_2026-07-01.md",
        inventory,
        product_audit,
        price_wedge_audit,
        monthly,
    )

    print("Wrote Wind price/control audit outputs.")
    print(f"Inventory rows: {len(inventory)}")
    print(f"Clean long rows: {len(clean_long)}")
    print(f"Monthly panel shape: {monthly.shape}")


if __name__ == "__main__":
    main()
